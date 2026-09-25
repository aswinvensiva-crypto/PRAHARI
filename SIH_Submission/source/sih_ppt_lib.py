"""Shared helpers for filling the official SIH idea-presentation template (see build_*_ppt.py)."""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


SRC = Path(__file__).resolve().parent
TEMPLATE = SRC / "SIH2025-IDEA-Presentation-Format.pptx"
ICONS = SRC / "icons"

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x00, 0x70, 0xC0)
ORANGE = RGBColor(0xE8, 0x77, 0x22)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xD6, 0x45, 0x45)
AMBER = RGBColor(0xE0, 0xA8, 0x00)
DARK = RGBColor(0x2B, 0x2B, 0x2B)
INK = RGBColor(0x1C, 0x27, 0x33)
MUTED = RGBColor(0x59, 0x65, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
T_BLUE = RGBColor(0xEA, 0xF1, 0xFB)
T_ORANGE = RGBColor(0xFD, 0xF0, 0xE2)
T_GREEN = RGBColor(0xE6, 0xF4, 0xEC)
T_GREY = RGBColor(0xF3, 0xF5, 0xF8)
LINE = RGBColor(0xC9, 0xD3, 0xDD)
FONT = "Arial"


# ---------------------------------------------------------------- helpers
def _bullet(p, char="•", color=None):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.17)))
    pPr.set("indent", str(-Inches(0.17)))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if color is not None:
        buClr = pPr.makeelement(qn("a:buClr"), {})
        clr = buClr.makeelement(qn("a:srgbClr"), {"val": str(color)})
        buClr.append(clr)
        pPr.append(buClr)
    buFont = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    pPr.append(buFont)
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char}))


def fill_tf(tf, paras, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, space_after=2, line=None):
    """paras: list of str | dict(runs=[(text, opts)], bullet=bool, size, color, bold, align, after)."""
    tf.word_wrap = True
    first = True
    for item in paras:
        if isinstance(item, str):
            item = {"runs": [(item, {})]}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = item.get("align", align)
        p.space_after = Pt(item.get("after", space_after))
        if line or item.get("line"):
            p.line_spacing = item.get("line", line)
        if item.get("bullet"):
            _bullet(p, color=item.get("bullet_color", ORANGE))
        for text, o in item["runs"]:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", item.get("size", size)))
            f.bold = o.get("bold", item.get("bold", bold))
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", item.get("color", color))
            if o.get("link"):
                r.hyperlink.address = o["link"]
    return tf


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.03, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    fill_tf(tf, paras, **kw)
    return tb


def shape(slide, kind, x, y, w, h, fill=None, line=None, paras=None, anchor=MSO_ANCHOR.MIDDLE, margin=0.08,
          radius=None, line_w=1.0, **kw):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    tf = s.text_frame
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    if paras:
        fill_tf(tf, paras, **kw)
    return s


def pointer(slide, x, y, w, label, h=0.36, size=13):
    """Template pointer text, kept verbatim, shown as a section label."""
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=NAVY, radius=0.5, margin=0.14,
                 paras=[{"runs": [("❖  ", {"color": ORANGE}), (label, {})]}], size=size, bold=True, color=WHITE)


def icon(slide, name, x, y, d, color):
    shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=color)
    pad = d * 0.24
    slide.shapes.add_picture(str(ICONS / f"{name}.png"), Inches(x + pad), Inches(y + pad), Inches(d - 2 * pad), Inches(d - 2 * pad))


def arrow(slide, x, y, w, h, color=MUTED, kind=MSO_SHAPE.RIGHT_ARROW):
    return shape(slide, kind, x, y, w, h, fill=color)


def table(slide, x, y, w, col_w, rows, header_fill=NAVY, size=10.5, row_h=0.3, first_col_bold=True, zebra=True):
    gf = slide.shapes.add_table(len(rows), len(col_w), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = T_GREY if (zebra and r % 2 == 0) else WHITE
            tf = cell.text_frame
            tf.paragraphs[0].text = ""
            runs = val if isinstance(val, list) else [(val, {})]
            fill_tf(tf, [{"runs": runs}], size=size,
                    bold=(r == 0) or (c == 0 and first_col_bold),
                    color=WHITE if r == 0 else (NAVY if c == 0 and first_col_bold else INK))
    return tbl


def remove(shape_):
    el = shape_._element
    el.getparent().remove(el)


def find(slide, name_prefix):
    return next(s for s in slide.shapes if s.name.startswith(name_prefix))


def set_logo(slide, logo_text, size=13):
    oval = next(s for s in slide.shapes if s.name.startswith("Oval"))
    p = oval.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = logo_text
    runs[0].font.bold = True
    runs[0].font.size = Pt(size)
    oval.text_frame.word_wrap = False
    oval.text_frame.margin_left = oval.text_frame.margin_right = 0
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for br in p._p.findall(qn("a:br")):
        p._p.remove(br)
    for extra in oval.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def set_title(slide, title, size):
    ph = slide.shapes.title
    ph.left, ph.width = Inches(1.85), Inches(8.8)
    p = ph.text_frame.paragraphs[0]
    keep = p.runs[0]  # keep the template's run formatting (font, bold, colour)
    for child in list(p._p):
        if child.tag not in (qn("a:pPr"), qn("a:endParaRPr")) and child is not keep._r:
            p._p.remove(child)
    keep.text = title
    keep.font.size = Pt(size)



def open_template():
    """Open the template and apply rule 1 (max six slides): drop the "Important Instructions" slide,
    which the template's own note says may be deleted."""
    prs = Presentation(str(TEMPLATE))
    sld_ids = prs.slides._sldIdLst
    last = sld_ids[-1]
    prs.part.drop_rel(last.rId)
    sld_ids.remove(last)
    slides = list(prs.slides)
    assert len(slides) == 6
    return prs, slides


def pointers_of(slide, remove_box=True):
    """Return the template's idea-detail pointer texts (verbatim) from the slide's pointer text box."""
    tb = find(slide, "TextBox 8")
    texts = [p.text.strip() for p in tb.text_frame.paragraphs if p.text.strip()]
    if remove_box:
        remove(tb)
    return texts


def fill_title_slide(slide, team, idea_runs):
    tb = find(slide, "TextBox 9")
    values = [team["ps_id"], team["ps_title"], team["theme"], team["category"], team["team_id"], team["team_name"]]
    paras = [p for p in tb.text_frame.paragraphs if p.runs]
    for p, val in zip(paras, values):
        label = p.runs[0]
        label.text = "PS Category- " if label.text.startswith("PS Category") else label.text.rstrip() + " "
        label.font.size = Pt(18)
        p.alignment = PP_ALIGN.LEFT
        p._p.append(copy.deepcopy(label._r))
        p.runs[-1].text = val
        p.runs[-1].font.bold = False
        p.runs[-1].font.color.rgb = NAVY
    text(slide, 0.36, 2.2, 6.3, 0.7, [{"runs": idea_runs}], size=13)
